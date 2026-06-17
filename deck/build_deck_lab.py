#!/usr/bin/env python3
"""
Build the MX17 resistive-Micromegas DETECTOR-PERFORMANCE deck, aimed at the lab
that produced the detectors (cosmic test bench at CEA Saclay).

Focus: how the detector performs — efficiency, spatial resolution, angular
resolution (µ-TPC), and that it reconstructs tracks both on the cosmic bench and
in the n_TOF beam.  Short main section, large backup.  The gamma-flash / DAQ
saturation story is deliberately left out (covered in the separate status deck).
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# ---- palette ----------------------------------------------------------------
NAVY   = RGBColor(0x14, 0x2C, 0x52)
BLUE   = RGBColor(0x2E, 0x6D, 0xB4)
ORANGE = RGBColor(0xE0, 0x78, 0x20)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)
GREY   = RGBColor(0x55, 0x55, 0x55)
LGREY  = RGBColor(0xEE, 0xEE, 0xEE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x22, 0x22, 0x22)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set(tf_para, text, size, color=DARK, bold=False, italic=False, align=None, font="Calibri"):
    tf_para.text = text
    r = tf_para.runs[0]
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    if align is not None:
        tf_para.alignment = align


def add_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s


def add_band(s, color=NAVY, h=Inches(0.16), y=0):
    b = s.shapes.add_shape(1, 0, y, SW, h)
    b.fill.solid(); b.fill.fore_color.rgb = color; b.line.fill.background()
    b.shadow.inherit = False
    return b


def title_box(s, title, kicker=None):
    add_band(s, NAVY, Inches(0.14), 0)
    if kicker:
        k = s.shapes.add_textbox(Inches(0.55), Inches(0.20), Inches(12), Inches(0.35))
        _set(k.text_frame.paragraphs[0], kicker.upper(), 12.5, ORANGE, bold=True)
    ty = Inches(0.52) if kicker else Inches(0.30)
    t = s.shapes.add_textbox(Inches(0.55), ty, Inches(12.3), Inches(0.9))
    tf = t.text_frame; tf.word_wrap = True
    _set(tf.paragraphs[0], title, 27, NAVY, bold=True)
    return t


def footer(s, n):
    f = s.shapes.add_textbox(Inches(11.3), Inches(7.06), Inches(1.9), Inches(0.32))
    _set(f.text_frame.paragraphs[0], f"MX17 detector  ·  {n}", 9, GREY, align=PP_ALIGN.RIGHT)


def bullets(s, items, x, y, w, h, size=15, gap=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0; txt = it; color = DARK; bold = False
        if isinstance(it, tuple):
            txt, lvl = it[0], it[1]
            if len(it) > 2: color = it[2]
            if len(it) > 3: bold = it[3]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bullet = ("•  " if lvl == 0 else "–  ") if txt else ""
        _set(p, bullet + txt, size - lvl*1.0, color, bold=bold)
        p.level = lvl
        p.space_after = Pt(gap)
    return tb


def fit_image(s, path, x, y, w, h, caption=None, frame=True):
    """Place image fit within box (w,h) preserving aspect, centred."""
    iw, ih = Image.open(path).size
    ar = iw / ih
    bw, bh = w, h
    if bw / bh > ar:
        nh = bh; nw = int(bh * ar)
    else:
        nw = bw; nh = int(bw / ar)
    nx = x + (bw - nw) // 2
    ny = y + (bh - nh) // 2
    if frame:
        bg = s.shapes.add_shape(1, nx-Emu(9525), ny-Emu(9525), nw+Emu(19050), nh+Emu(19050))
        bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
        bg.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC); bg.line.width = Pt(0.75)
        bg.shadow.inherit = False
    pic = s.shapes.add_picture(path, nx, ny, nw, nh)
    if caption:
        c = s.shapes.add_textbox(x, y + h, w, Inches(0.3))
        _set(c.text_frame.paragraphs[0], caption, 10.5, GREY, italic=True, align=PP_ALIGN.CENTER)
    return pic


def A(name):  # asset path
    return os.path.join(ASSETS, name)


def callout(s, text, x, y, w, h, fill=RGBColor(0xFB, 0xF3, 0xD8), line=ORANGE, size=14, bold=False, color=DARK):
    box = s.shapes.add_shape(5, x, y, w, h)  # rounded rect
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.color.rgb = line; box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    parts = text.split("\n")
    for i, pt in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, pt, size, color, bold=bold)
    return box


def section_divider(title, subtitle=None):
    s = add_slide(NAVY)
    add_band(s, ORANGE, Inches(0.16), Inches(3.0))
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.0))
    _set(t.text_frame.paragraphs[0], title, 36, WHITE, bold=True)
    if subtitle:
        st = s.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(11.7), Inches(0.8))
        _set(st.text_frame.paragraphs[0], subtitle, 17, RGBColor(0xCF, 0xDB, 0xEC))
    return s


def table(s, data, x, y, w, h, header=True, fs=11, col_w=None):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, x, y, w, h).table
    if col_w:
        for j, cw in enumerate(col_w):
            gt.columns[j].width = cw
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = gt.cell(i, j)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            if i == 0 and header:
                _set(p, str(val), fs, WHITE, bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            else:
                _set(p, str(val), fs, DARK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if (i % 2) else LGREY
    return gt


GAS_TAG = "Argon/Isobutane 95/5  ·  30 mm drift  ·  aluminium frame  ·  cosmic bench @ CEA Saclay"

# =============================================================================
#  TITLE
# =============================================================================
s = add_slide(NAVY)
add_band(s, ORANGE, Inches(0.18), Inches(2.55))
t = s.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.7), Inches(1.5))
_set(t.text_frame.paragraphs[0], "MX17 Resistive Micromegas — Detector Performance", 36, WHITE, bold=True)
sub = s.shapes.add_textbox(Inches(0.8), Inches(2.75), Inches(11.7), Inches(1.2))
tf = sub.text_frame; tf.word_wrap = True
_set(tf.paragraphs[0], "Cosmic-bench characterisation at CEA Saclay and first reconstructed tracks at n_TOF EAR2",
     20, RGBColor(0xCF, 0xDB, 0xEC))
p2 = tf.add_paragraph()
_set(p2, "Prototype Detector 0 — how the detector you built performs", 15, RGBColor(0xAE, 0xC3, 0xDC), italic=True)
meta = s.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.8))
_set(meta.text_frame.paragraphs[0], "n_TOF X17  ·  prepared for the detector lab  ·  June 2026   [add recipient / affiliation]",
     13, RGBColor(0x9A, 0xAE, 0xC8))

# =============================================================================
#  MAIN — short
# =============================================================================
# 1. Physics context — one slide
s = add_slide(); title_box(s, "What these detectors are for", kicker="Context · one slide")
bullets(s, [
    ("The ATOMKI experiments report an excess of e⁺e⁻ pairs peaking near a 120° opening", 0),
    ("angle — possibly a new ~17 MeV boson, the “X17”.", 1),
    ("n_TOF EAR2 can make an independent, complementary measurement using a ³He target:", 0),
    ("n + ³He → ⁴He* → ⁴He + X17 → e⁺e⁻.", 1),
    ("The observable is the e⁺e⁻ opening angle → we must reconstruct the two tracks and the", 0, NAVY, True),
    ("angle between them, event by event.", 1, NAVY, True),
    ("Our tool: four 40×40 cm resistive-strip Micromegas in µ-TPC mode (30 mm drift),", 0),
    ("giving full 3-D track angles around the target.", 1),
    ("→ The detector requirement is good efficiency to MIPs and good ANGULAR resolution.", 0, GREEN, True),
], Inches(0.55), Inches(1.6), Inches(7.4), Inches(5.2), size=15.5)
fit_image(s, A("fig_atomki_peak.png"), Inches(8.2), Inches(2.0), Inches(4.9), Inches(3.0),
          caption="ATOMKI: e⁺e⁻ excess near ~120° — the signature we aim to reconstruct.")
callout(s, "The rest of this deck is about how the detector performs —\nnot the physics. Efficiency, spatial & angular resolution, tracking.",
        Inches(8.2), Inches(5.4), Inches(4.9), Inches(1.2), size=13.5, bold=True, color=NAVY)
footer(s, "Context")

# 2. How we characterised it — cosmic bench
s = add_slide(); title_box(s, "Characterised on the cosmic test bench at CEA Saclay", kicker="Method")
fit_image(s, A("cb_setup_photo.png"), Inches(0.5), Inches(1.55), Inches(4.4), Inches(5.5),
          caption="Cosmic bench (side view).")
bullets(s, [
    ("Prototype “Detector 0”: 40×40 cm resistive-strip Micromegas, 30 mm drift gap,", 0),
    ("aluminium frame, Argon/Isobutane 95/5, DREAM readout.", 1),
    ("Mounted in a cosmic-muon telescope: 4 reference tracking detectors + 2 trigger", 0),
    ("scintillators provide an external reference track (the M3 tracker).", 1),
    ("Cosmic muons (MIPs) cross all layers; we compare the MX17 reconstruction to the", 0),
    ("reference track to extract efficiency, spatial resolution and angular resolution.", 1),
    ("µ-TPC mode: strip position + drift time → the 3rd coordinate, so a single detector", 0, NAVY, True),
    ("reconstructs a 3-D track segment and its angle.", 1, NAVY, True),
    ("Reference tracker resolution ≈ 500 µm — it sets the floor on what we can measure.", 0, GREY, True),
], Inches(5.3), Inches(1.7), Inches(7.8), Inches(5.2), size=14.5)
footer(s, "Method")

# 3. Efficiency
s = add_slide(); title_box(s, "The detector is efficient", kicker="Performance · efficiency")
fit_image(s, A("cb_efficiency.png"), Inches(0.4), Inches(1.7), Inches(8.7), Inches(3.9),
          caption="Per-bin efficiency = MX17 hits / reference tracks (left); reference-track density / exposure (right). 42,923 reference tracks.")
bullets(s, [
    ("Across most of the active area the detector is at or near", 0),
    ("full efficiency to cosmic MIPs.", 1),
    ("The two understood dead features:", 0, RED, True),
    ("a band of broken resistive strips (no HV there);", 1, RED, False),
    ("the M4 connector edge not making contact.", 1, RED, False),
    ("Both causes are identified and fixable in the next", 0, GREEN, True),
    ("detector iteration — not intrinsic to the design.", 1, GREEN, False),
], Inches(9.3), Inches(1.7), Inches(3.8), Inches(5.0), size=14)
g = s.shapes.add_textbox(Inches(0.4), Inches(6.55), Inches(12.5), Inches(0.4))
_set(g.text_frame.paragraphs[0], GAS_TAG, 11, GREY, italic=True)
footer(s, "Performance")

# 4. Spatial resolution
s = add_slide(); title_box(s, "Spatial resolution ~500 µm", kicker="Performance · spatial")
fit_image(s, A("cb_spatial_res.png"), Inches(0.4), Inches(1.7), Inches(12.5), Inches(3.8),
          caption="Sliding-window residual resolution maps in X and Y, and event statistics per bin.")
bullets(s, [
    ("Residual width (MX17 − reference) is ~500 µm over most of the detector.", 0, NAVY, True),
    ("The reference tracker itself only has ~500 µm resolution — so the measurement is", 0),
    ("resolution-limited by the bench: the true detector resolution is at least this good,", 1),
    ("probably better.", 1),
], Inches(0.55), Inches(5.7), Inches(12.4), Inches(1.5), size=14.5)
g = s.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(12.5), Inches(0.35))
_set(g.text_frame.paragraphs[0], GAS_TAG, 10.5, GREY, italic=True)
footer(s, "Performance")

# 5. Angular resolution — the headline number
s = add_slide(); title_box(s, "Angular resolution ~2° in µ-TPC mode", kicker="Performance · angular")
fit_image(s, A("cb_ang_res.png"), Inches(0.45), Inches(1.75), Inches(8.0), Inches(4.2),
          caption="µ-TPC angular resolution σ vs reference track angle (X and Y projections).")
bullets(s, [
    ("Angular resolution ≈ 2° across most of the angular range.", 0, GREEN, True),
    ("Worse only near 0° (tracks normal to the detector), where", 0),
    ("the µ-TPC has no longitudinal lever arm — expected.", 1),
    ("This is the key number for the physics: the X17 signature", 0, NAVY, True),
    ("is a ~120° e⁺e⁻ pair, so few-degree per-track resolution", 1, NAVY, True),
    ("is what we need.", 1, NAVY, True),
    ("Reconstruction algorithm not yet fully optimised — likely", 0),
    ("room to improve further.", 1),
], Inches(8.6), Inches(1.75), Inches(4.5), Inches(5.0), size=14)
g = s.shapes.add_textbox(Inches(0.4), Inches(6.6), Inches(12.5), Inches(0.35))
_set(g.text_frame.paragraphs[0], GAS_TAG, 10.5, GREY, italic=True)
footer(s, "Performance")

# 6. Works in the beam too — n_TOF event display
s = add_slide(); title_box(s, "…and it reconstructs tracks in the n_TOF beam", kicker="Performance · in beam")
fit_image(s, A("event_3d_geo.png"), Inches(0.3), Inches(1.5), Inches(7.9), Inches(5.6),
          caption="A reconstructed beam event in the n_TOF EAR2 setup (He/Ethane, 16 mm drift).")
bullets(s, [
    ("The same µ-TPC reconstruction works on real beam data at", 0, NAVY, True),
    ("n_TOF EAR2, not just on the bench.", 1, NAVY, True),
    ("Single, well-separated tracks are cleanly reconstructed in", 0),
    ("the low-occupancy (thermal-neutron) window.", 1),
    ("Strip hits + drift time give the 3-D track (green); it crosses", 0),
    ("the Micromegas and continues toward the trigger scintillators.", 1),
    ("Confirms the detector behaves in the real experiment as it", 0, GREEN, True),
    ("did on the Saclay bench.", 1, GREEN, False),
], Inches(8.4), Inches(1.7), Inches(4.7), Inches(5.0), size=14)
footer(s, "In beam")

# 7. Summary
s = add_slide(); title_box(s, "Summary — the detector performs to spec", kicker="Bottom line")
bullets(s, [
    ("Efficient to MIPs across the active area; the only dead regions have understood,", 0, GREEN, True),
    ("fixable causes (broken resistive strips, M4 connector contact).", 1, GREEN, False),
    ("Spatial resolution ~500 µm — at the limit of what the reference bench can measure.", 0, NAVY, True),
    ("Angular resolution ~2° in µ-TPC mode — the number that matters for the X17 e⁺e⁻ angle.", 0, NAVY, True),
    ("Full 3-D track reconstruction demonstrated both on the cosmic bench and in the n_TOF beam.", 0, NAVY, True),
    ("Reconstruction software not yet fully optimised → headroom on resolution.", 0),
    ("Clear, concrete improvements identified for the production batch (strip break, connector).", 0, GREEN, True),
    ("Net: the detector you built does what the measurement needs.", 0, GREEN, True),
], Inches(0.6), Inches(1.8), Inches(12.4), Inches(5.0), size=16.5)
footer(s, "Summary")

# =============================================================================
#  BACKUP
# =============================================================================
section_divider("Backup", "Reconstruction detail, more event displays, specs & next iteration")


def img_slide(title, kicker, asset, caption=None, bullet_items=None, tag=None):
    s = add_slide(); title_box(s, title, kicker=kicker)
    if bullet_items:
        fit_image(s, A(asset), Inches(0.45), Inches(1.6), Inches(8.0), Inches(5.0), caption=caption)
        bullets(s, bullet_items, Inches(8.6), Inches(1.7), Inches(4.6), Inches(5.0), size=13)
    else:
        fit_image(s, A(asset), Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.0), caption=caption)
    if tag:
        g = s.shapes.add_textbox(Inches(0.4), Inches(6.95), Inches(12.5), Inches(0.35))
        _set(g.text_frame.paragraphs[0], tag, 10.5, GREY, italic=True)
    footer(s, "Backup")
    return s

# B1: µ-TPC reconstruction of one cosmic event (3-D)
img_slide("µ-TPC reconstruction — single cosmic event (3-D)", "Backup · Cosmic bench",
          "cb_mutpc_3d.png",
          caption="Event 33843. X-strip and Y-strip hits placed in the 30 mm drift window; the MX17 µ-TPC track sits on top of the reference track (M3) — good agreement.",
          tag=GAS_TAG)
# B2: the slope fits behind the µ-TPC angle
img_slide("µ-TPC reconstruction — drift-time vs position fits", "Backup · Cosmic bench",
          "cb_mutpc_fit.png",
          caption="Same event: drift time vs X and Y strip position. The slope of each line is the projected track angle; the intercept locates the mesh crossing. This is the µ-TPC measurement.",
          tag=GAS_TAG)
# B3: angle correlation reference vs detector
img_slide("Reconstructed vs reference track angle", "Backup · Cosmic bench",
          "cb_angle_corr.png",
          caption="MX17 µ-TPC angle vs reference-tracker angle (X and Y), v_drift = 35 µm/ns. Clear correlation along y = x; the offset is under study (possible reconstruction/algorithmic effect).",
          tag=GAS_TAG)
# B4: detector is alive — single muon raw
s = add_slide(); title_box(s, "Detector is alive — a single cosmic muon", kicker="Backup · Cosmic bench")
fit_image(s, A("cb_muon_2dmap.png"), Inches(0.5), Inches(1.65), Inches(6.0), Inches(4.6),
          caption="2-D hit map, single cosmic muon (Det-0 daytime run).")
fit_image(s, A("cb_muon_yt.png"), Inches(6.8), Inches(1.65), Inches(6.0), Inches(4.6),
          caption="Y position vs drift time for the same muon — the straight µ-TPC line; marker size = pulse amplitude.")
footer(s, "Backup")

# B5–B8: event_100 displays (n_TOF beam)
img_slide("Beam event display — Micromegas strip hits (3-D)", "Backup · n_TOF event display",
          "event_3d_mm.png",
          caption="Run 51, Event 100. FEU 1 X-strips (7 hits), FEU 2 Y-strips (9 hits) placed in 3-D from strip position and drift time (v_drift = 0.0320 mm/ns, 16 mm gap); green = fitted µ-TPC track.")
img_slide("Beam event display — hit-vs-drift fits & raw waveforms", "Backup · n_TOF event display",
          "event_2d.png",
          caption="Same event. Top: X/Y hit position vs drift time (slope = µ-TPC angle). Bottom: reconstructed track position and the raw X/Y strip waveforms.")
img_slide("Beam event display — X–Y detector-plane projection", "Backup · n_TOF event display",
          "event_xy.png",
          caption="Same event projected on the detector plane: X/Y strip hits, MM active area, scintillators and capsule, with the track crossing points at each plane.")
img_slide("Beam event display — cinematic render", "Backup · n_TOF event display",
          "event_cinematic.png",
          caption="3-D render of the reconstructed track through the resistive Micromegas. A rotating animation is available (event_100_3d_geo_rotating.gif).")
# B9: isolated thermal track in beam
img_slide("Isolated track in the n_TOF beam (thermal region)", "Backup · n_TOF",
          "fig_thermal_track.png",
          caption="He/Ethane, 30 mm drift, µ-TPC mode: a clean, well-separated track in the thermal-neutron window — naively exactly a single-track event (high dE/dx; particle ID uncertain).")

# B10: detector spec table
s = add_slide(); title_box(s, "Prototype Detector 0 — specifications", kicker="Backup · Specs")
table(s, [
    ["Parameter", "Value"],
    ["Type", "Resistive-strip bulk Micromegas (MX17)"],
    ["Active area", "40 × 40 cm"],
    ["Drift gap", "30 mm (µ-TPC mode)"],
    ["Frame", "Aluminium (carbon frame also tested)"],
    ["Readout", "X & Y strips, DREAM front-end (FEU)"],
    ["Bench gas", "Argon/Isobutane 95/5"],
    ["Reference tracker", "M3 cosmic telescope, ~500 µm resolution"],
    ["Cosmic statistics", "~43k reference tracks for efficiency/resolution maps"],
], Inches(0.6), Inches(1.7), Inches(9.0), Inches(4.8), fs=14, col_w=[Inches(3.6), Inches(5.4)])
bullets(s, [
    ("Performance", 0, NAVY, True),
    ("Efficiency: near-full over active area", 0),
    ("Spatial: ~500 µm (bench-limited)", 0),
    ("Angular: ~2° (µ-TPC)", 0),
], Inches(9.9), Inches(1.9), Inches(3.2), Inches(3.0), size=14)
footer(s, "Backup")

# B11: known issues & next iteration
s = add_slide(); title_box(s, "Known issues & fixes for the next iteration", kicker="Backup · Outlook")
table(s, [
    ["Observation", "Cause", "Fix for production batch"],
    ["Dead band in efficiency map", "Broken resistive strips (no HV in region)",
     "Improve resistive-strip QC / handling in fabrication"],
    ["Dead edge in efficiency map", "M4 connector edge not making contact",
     "Connector seating / mechanical revision"],
    ["σ(angle) larger near 0°", "No µ-TPC lever arm at normal incidence",
     "Intrinsic to µ-TPC; mitigated by geometry around target"],
    ["Angle-correlation offset", "Possible reconstruction/algorithm effect",
     "Reconstruction software optimisation (in progress)"],
], Inches(0.55), Inches(1.7), Inches(12.25), Inches(3.4), fs=13,
   col_w=[Inches(3.4), Inches(4.2), Inches(4.65)])
callout(s, "Bottom line: every observed limitation is either understood and fixable in fabrication, "
           "or a known µ-TPC property — none point to a problem with the detector concept. "
           "Four production detectors are planned for the July n_TOF run.",
        Inches(0.55), Inches(5.3), Inches(12.25), Inches(1.4), size=14.5, bold=True, color=NAVY)
footer(s, "Backup")

out = os.path.join(HERE, "MX17_Detector_Performance.pptx")
prs.save(out)
print("Saved", out, "with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
